"""Local, serial DeckFoundry HTTP frontend. No third-party web dependencies."""
import html
import base64
import io
import zipfile
from lxml import html as dom
from pptx import Presentation
from PIL import Image
import ipaddress
import http.cookies
import json
import itertools
import signal
import os
from pathlib import Path
import re
import secrets
import subprocess
import sys
import time
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib import request, parse

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
import project

LIMIT = 65536
SOURCE_LIMIT = 2 * 1024 * 1024
ID = re.compile(r'[A-Za-z0-9_-]{1,80}\Z')
VERSION = re.compile(r'v[0-9]{4,8}\Z')
# Only leaf text elements are editable. Markup, attributes and layout stay intact.
SLOT = re.compile(r'<(h1|h2|h3|p|li|aside|div)\b([^<>]*)>([^<>]*)</\1\s*>', re.I)
STATIC = Path(__file__).parent


def normalize_origin(value):
    parsed = parse.urlsplit(value)
    if (parsed.scheme not in ('http', 'https') or not parsed.netloc or parsed.username or parsed.password
            or parsed.path not in ('', '/') or parsed.query or parsed.fragment):
        raise ValueError('Invalid origin')
    parsed.port  # Validate a supplied port.
    return parsed.scheme.lower() + '://' + parsed.netloc.lower()


def allowed_origins():
    configured = os.environ.get('DECKFOUNDRY_ALLOWED_ORIGINS', '').strip()
    if configured:
        values = configured.split(',')
    else:
        port = int(os.environ.get('DECKFOUNDRY_PORT', '8080'))
        values = [f'http://127.0.0.1:{port}', f'http://localhost:{port}']
    return {normalize_origin(value.strip()) for value in values if value.strip()}


def bounded_text(value, maximum=4000):
    if not isinstance(value, str) or len(value) > maximum or any(ord(c) < 32 and c not in '\n\t' for c in value):
        raise ValueError('Invalid text')
    return value


BLANK = '<section class="slide" style="background:#ffffff"><h1>Untitled presentation</h1><p>Click to tell your story</p></section>'
SHELL = '<html><head><style>.stage{width:1440px;height:810px}.slide{position:absolute;inset:0;padding:80px;box-sizing:border-box;font-family:Arial}.slide h1{font-size:52px}.slide p{font-size:28px;white-space:pre-wrap}</style></head><body><div class="stage">%s</div></body></html>'


def safe_style(value):
    result = {}
    for key, val in re.findall(r'([a-z-]+)\s*:\s*([^;]+)', value or ''):
        val = val.strip()
        if key in ('color','background') and re.fullmatch(r'#[0-9a-fA-F]{6}', val):
            result[key] = val
        elif key == 'font-size' and re.fullmatch(r'(?:[1-9][0-9]?|1[0-5][0-9]|160)px', val):
            result[key] = val
        elif key == 'text-align' and val in ('left','center','right'):
            result[key] = val
        elif key in ('left','top','width','height') and re.fullmatch(r'[0-9]{1,4}(?:\.[0-9]+)?px', val):
            result[key] = val
    if 'left' in result or 'top' in result:
        result['position'] = 'absolute'
    return result


def checked_image(raw):
    try:
        if len(raw)>1024*1024:
            raise ValueError('Image size limit')
        with Image.open(io.BytesIO(raw)) as image:
            if image.format not in ('PNG','JPEG') or image.width*image.height>16000000:
                raise ValueError('Image dimensions or format')
            image.verify()
    except Exception as exc:
        raise ValueError('Import contains an invalid or oversized image') from exc
    return raw


def import_source(filename, raw):
    if not isinstance(filename,str) or not re.fullmatch(r'[A-Za-z0-9 _.-]{1,100}',filename) or filename.startswith('.'):
        raise ValueError('Use a simple filename without paths')
    if not raw or len(raw)>8*1024*1024:
        raise ValueError('Upload limit is 8 MiB')
    sections=[]
    suffix=Path(filename).suffix.lower()
    if suffix in ('.html','.htm'):
        if len(raw)>SOURCE_LIMIT:
            raise ValueError('HTML limit is 2 MiB')
        text=raw.decode('utf-8')
        if '<!entity' in text.lower() or '<!doctype' in text.lower() and 'system' in text.lower():
            raise ValueError('External declarations unsupported')
        try:
            tree=dom.fromstring(text)
        except (ValueError, dom.etree.ParserError) as exc:
            raise ValueError('Malformed HTML file') from exc
        for bad in tree.xpath('//script|//style|//iframe|//object|//embed|//svg|//form'):
            bad.drop_tree()
        for node in tree.xpath('//section[contains(concat(" ",normalize-space(@class)," ")," slide ")]'):
            children=[]
            elements=node.xpath('.//h1|.//h2|.//h3|.//p|.//li')
            if len(elements)>80: raise ValueError('Import has more than 80 elements on a slide')
            for el in elements:
                content=bounded_text(''.join(el.itertext()))
                style=';'.join(k+':'+v for k,v in safe_style(el.get('style')).items())
                children.append('<'+el.tag+' style="'+style+'">'+html.escape(content,quote=False)+'</'+el.tag+'>')
            images=node.xpath('.//img')
            if len(images)+len(elements)>80: raise ValueError('Import has more than 80 elements on a slide')
            for image in images:
                src=image.get('src','')
                if len(src)<1400000 and re.fullmatch(r'data:image/(?:png|jpeg);base64,[A-Za-z0-9+/=]+',src):
                    checked_image(base64.b64decode(src.split(',',1)[1],validate=True))
                    style=';'.join(k+':'+v for k,v in safe_style(image.get('style')).items())
                    children.append('<img style="'+style+'" src="'+src+'">')
            section_style=';'.join(k+':'+v for k,v in safe_style(node.get('style')).items())
            sections.append('<section class="slide" style="'+section_style+'">'+''.join(children)+'</section>')
    elif suffix=='.pptx':
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries=archive.infolist()
                if len(entries)>2000 or sum(e.file_size for e in entries)>32*1024*1024:
                    raise ValueError('Expanded package exceeds limits')
                for entry in entries:
                    if entry.filename.startswith('/') or '..' in Path(entry.filename).parts or entry.flag_bits&1:
                        raise ValueError('Unsafe package path')
                    if entry.filename.endswith(('.xml','.rels')):
                        part=archive.read(entry).upper()
                        if b'<!DOCTYPE' in part or b'<!ENTITY' in part:
                            raise ValueError('External declarations unsupported')
            deck=Presentation(io.BytesIO(raw))
            if not 1<=len(deck.slides)<=200:
                raise ValueError('Slide limit exceeded')
            for slide in deck.slides:
                children=[]
                if len(slide.shapes)>80: raise ValueError('Import has more than 80 elements on a slide')
                for shape in slide.shapes:
                    style=f'position:absolute;left:{shape.left/deck.slide_width*1440:.2f}px;top:{shape.top/deck.slide_height*810:.2f}px;width:{shape.width/deck.slide_width*1440:.2f}px;height:{shape.height/deck.slide_height*810:.2f}px'
                    if shape.has_text_frame:
                        runs=[r for p in shape.text_frame.paragraphs for r in p.runs]
                        size=min(160,max(10,round(runs[0].font.size.pt*1.25))) if runs and runs[0].font.size else 28
                        color='#172b4d'
                        if runs:
                            try: color='#'+str(runs[0].font.color.rgb) if runs[0].font.color.rgb else color
                            except (AttributeError,ValueError): pass
                        align=str(shape.text_frame.paragraphs[0].alignment)
                        alignment='center' if 'CENTER' in align else 'right' if 'RIGHT' in align else 'left'
                        children.append('<p style="'+style+f';font-size:{size}px;color:{color};text-align:{alignment}">'+html.escape(bounded_text(shape.text),quote=False)+'</p>')
                    elif shape.shape_type==13 and shape.image.content_type in ('image/png','image/jpeg') and len(shape.image.blob)<1024*1024:
                        children.append('<img style="'+style+'" src="data:'+shape.image.content_type+';base64,'+base64.b64encode(checked_image(shape.image.blob)).decode()+'">')
                sections.append('<section class="slide">'+''.join(children)+'</section>')
        except (zipfile.BadZipFile,KeyError,OSError,dom.etree.XMLSyntaxError) as exc:
            raise ValueError('Malformed PowerPoint file') from exc
    else:
        raise ValueError('Choose an HTML or PPTX file')
    source=SHELL % ''.join(sections)
    if not 1<=len(sections)<=200 or len(source.encode())>SOURCE_LIMIT:
        raise ValueError('Import needs 1–200 slides within 2 MiB editable source')
    return source


class App:
    def __init__(self, root):
        self.root = Path(root).absolute()
        if self.root.is_symlink() or self.root.resolve() != self.root:
            raise ValueError('Projects root must not contain symlinks')
        self.sessions = {}

    def path(self, *parts):
        p = self.root
        for part in parts:
            if not part or part in ('.', '..') or '/' in part or '\\' in part:
                raise ValueError('Invalid path')
            p = p / part
            if p.is_symlink():
                raise ValueError('Symlink rejected')
        if not p.resolve().is_relative_to(self.root):
            raise ValueError('Path escape')
        return p

    def load(self, name):
        if not ID.fullmatch(name):
            raise ValueError('Invalid project')
        root = self.path(name)
        # Reject symlinks anywhere the existing engine could read/write, including old versions.
        count = 0
        for directory, dirs, files in os.walk(root, followlinks=False):
            for item in dirs + files:
                count += 1
                if count > 10000 or (Path(directory) / item).is_symlink():
                    raise ValueError('Unsafe or oversized project tree')
        state_file = self.path(name, 'project.json')
        if state_file.stat().st_size > LIMIT:
            raise ValueError('Oversized metadata')
        state = json.loads(state_file.read_text())
        if not isinstance(state, dict) or state.get('authority') != 'html' or not VERSION.fullmatch(state.get('current', '')):
            raise ValueError('Only existing HTML-source projects supported')
        folder = self.path(name, 'versions', state['current'])
        changes_file = folder / 'changes.json'
        if changes_file.stat().st_size > LIMIT:
            raise ValueError('Oversized revision metadata')
        changes = json.loads(changes_file.read_text())
        parent = changes.get('parent')
        if parent is not None and (not isinstance(parent, str) or not VERSION.fullmatch(parent)):
            raise ValueError('Invalid parent revision')
        if parent:
            self.path(name, 'versions', parent)
        source = folder / 'source.html'
        if source.stat().st_size > SOURCE_LIMIT:
            raise ValueError('Source too large')
        text = source.read_text()
        fragments = project.slides(text)
        if not 1 <= len(fragments) <= 200:
            raise ValueError('Slide limit exceeded')
        return root, state, fragments

    def view(self, name, slide):
        _, state, fragments = self.load(name)
        if type(slide) is not int or not 1 <= slide <= len(fragments):
            raise ValueError('Invalid slide')
        slots = []
        for match in SLOT.finditer(fragments[slide-1].group()):
            tag, attrs, raw = match.groups()
            if len(slots) >= 80 or len(raw) > 4000:
                raise ValueError('Slide fields exceed limit')
            slots.append({'kind': 'title' if tag.lower() == 'h1' else 'notes' if 'notes' in attrs else 'bullet' if tag.lower() == 'li' else 'text', 'old': raw, 'text': html.unescape(raw), 'style': safe_style(dom.fromstring(match.group()).get('style'))})
        return {'project': name, 'base': state['current'], 'slide': slide,
                'slides': [x['title'][:200] for x in project.index(''.join(m.group() for m in fragments))],
                'background': safe_style(dom.fromstring(fragments[slide-1].group()).get('style')).get('background','#ffffff'),
                'images': [{'src':e.get('src'),'style':safe_style(e.get('style'))} for e in dom.fromstring(fragments[slide-1].group()).xpath('.//img') if re.fullmatch(r'data:image/(?:png|jpeg);base64,[A-Za-z0-9+/=]+',e.get('src',''))],
                'fields': slots, 'unsupported': bool(state.get('brand'))}

    def templates(self):
        folder=self.path('_templates')
        if not folder.exists():
            return []
        return [{'name':p.name} for p in itertools.islice(folder.iterdir(),1000) if not p.is_symlink() and p.is_file() and re.fullmatch(r'[A-Za-z0-9_-]{1,80}\.html',p.name)]

    def create(self,data):
        if set(data)-{'name','template','filename','content','kind'} or data.get('kind','project') not in ('project','template'):
            raise ValueError('Invalid creation payload')
        name=data.get('name','')
        if not ID.fullmatch(name) or name.startswith('_'):
            raise ValueError('Name needs letters, numbers, hyphens or underscores')
        if data.get('kind')=='template' and not data.get('filename'):
            raise ValueError('Choose a file to add a template')
        source=SHELL % BLANK
        raw=None
        if data.get('template'):
            filename=data['template']
            if filename not in {t['name'] for t in self.templates()}:
                raise ValueError('Template unavailable')
            template=self.path('_templates',filename)
            if template.stat().st_size>SOURCE_LIMIT: raise ValueError('Template exceeds 2 MiB')
            source=import_source(filename,template.read_bytes())
        elif data.get('filename'):
            raw=base64.b64decode(data.get('content',''),validate=True)
            source=import_source(data['filename'],raw)
        if data.get('kind')=='template':
            self.path('_templates').mkdir(exist_ok=True)
            with self.path('_templates',name+'.html').open('x') as output:
                output.write(source)
            return {'template':name+'.html'}
        root=self.path(name)
        if root.exists():
            raise ValueError('This project name is already in use')
        root.mkdir()
        original=root/'import.html';original.write_text(source)
        project.init(root,original)
        if raw is not None:
            (root/('original'+Path(data['filename']).suffix.lower())).write_bytes(raw)
        return {'project':name}

    def edit(self,data):
        root,state,fragments=self.load(data['project'])
        if data.get('base')!=state['current']:
            raise Conflict()
        n=data.get('slide')
        if type(n) is not int or not 1<=n<=len(fragments):
            raise ValueError('Invalid slide')
        source=(project.revision(root)[1]/'source.html').read_text()
        parts=[m.group() for m in fragments]
        action=data.get('action')
        if action in ('add','duplicate'):
            parts.insert(n,BLANK if action=='add' else parts[n-1])
        elif action=='delete':
            if len(parts)==1: raise ValueError('Keep at least one slide')
            parts.pop(n-1)
        elif action in ('up','down'):
            target=n-2 if action=='up' else n
            if not 0<=target<len(parts): raise ValueError('Slide is already at the edge')
            parts[n-1],parts[target]=parts[target],parts[n-1]
        elif action=='format':
            edits=data.get('edits',[data])
            if not isinstance(edits,list) or not 1<=len(edits)<=81:
                raise ValueError('Invalid edits')
            for edit in edits:
                if not isinstance(edit,dict): raise ValueError('Invalid edit')
                styles=edit.get('styles',{})
                if not isinstance(styles,dict) or set(styles)-{'font-size','color','text-align','background'}:
                    raise ValueError('Invalid properties')
                if safe_style(';'.join(k+':'+str(v) for k,v in styles.items()))!=styles:
                    raise ValueError('Invalid property value')
                ix=edit.get('field')
                if ix is None:
                    if set(styles)!={'background'}: raise ValueError('Select an element')
                    node=dom.fromstring(parts[n-1])
                    current=dict(re.findall(r'([a-z-]+)\s*:\s*([^;]+)',node.get('style','')));current.update(styles)
                    node.set('style',';'.join(k+':'+v for k,v in current.items()));parts[n-1]=dom.tostring(node,encoding='unicode')
                else:
                    fields=list(SLOT.finditer(parts[n-1]))
                    if type(ix) is not int or not 0<=ix<len(fields): raise ValueError('Invalid element')
                    match=fields[ix];node=dom.fromstring(match.group())
                    current=dict(re.findall(r'([a-z-]+)\s*:\s*([^;]+)',node.get('style','')));current.update(styles)
                    node.set('style',';'.join(k+':'+v for k,v in current.items()))
                    if 'text' in edit: node.text=bounded_text(edit['text'])
                    parts[n-1]=parts[n-1][:match.start()]+dom.tostring(node,encoding='unicode')+parts[n-1][match.end():]
        else: raise ValueError('Unknown slide action')
        # Keep the original shell and replace slide ranges in place.
        for i,match in reversed(list(enumerate(fragments))):
            replacement=''.join(parts[i:]) if i==len(fragments)-1 else parts[i] if i<len(parts) else ''
            source=source[:match.start()]+replacement+source[match.end():]
        return project.source_revision(root,state['current'],source)

    def validate_ops(self, view, ops):
        if not isinstance(ops, list) or not 1 <= len(ops) <= 40:
            raise ValueError('Expected 1–40 operations')
        allowed = {f['old'] for f in view['fields']}
        _, _, fragments = self.load(view['project'])
        fragment = fragments[view['slide'] - 1].group()
        seen = set()
        for op in ops:
            if not isinstance(op, dict) or set(op) != {'slide', 'old', 'new'}:
                raise ValueError('Invalid operation')
            if type(op['slide']) is not int or op['slide'] != view['slide'] or op['old'] not in allowed or op['old'] in seen:
                raise ValueError('Operation must match an active slide text field')
            bounded_text(op['old']); bounded_text(op['new'])
            # Ops contain exact source literals, with new content canonically escaped.
            if html.escape(html.unescape(op['new']), quote=False) != op['new']:
                raise ValueError('New text must be HTML escaped')
            if fragment.count(op['old']) != 1:
                raise ValueError('Old literal must match exactly once')
            fragment = fragment.replace(op['old'], op['new'], 1)
            seen.add(op['old'])
        return ops

    def apply(self, data):
        if set(data) != {'project', 'slide', 'base', 'ops'}:
            raise ValueError('Invalid apply payload')
        view = self.view(data['project'], data['slide'])
        if data['base'] != view['base']:
            raise Conflict()
        ops = self.validate_ops(view, data['ops'])
        root, _, _ = self.load(data['project'])
        result = project.patch_revision(root, {'base': data['base'], 'message': 'Web slide edit', 'ops': ops})
        return {'current': result['current']}

    def chat(self, session, data):
        if set(data) != {'project', 'slide', 'base', 'message'}:
            raise ValueError('Invalid chat payload')
        bounded_text(data['message'], 2000)
        view = self.view(data['project'], data['slide'])
        if data['base'] != view['base']:
            raise Conflict()
        base = os.environ.get('HERMES_API_BASE_URL', '')
        key = os.environ.get('HERMES_API_KEY', '')
        url = parse.urlsplit(base)
        if url.scheme not in ('http', 'https') or not url.netloc or not url.hostname or url.username or url.password or url.query or url.fragment or not key:
            raise ValueError('Chat is not configured')
        if url.scheme == 'http':
            try:
                address = ipaddress.ip_address(url.hostname)
                local = address.is_loopback
                private = address.is_private
            except ValueError:
                local = url.hostname.lower() == 'localhost'
                private = False
            if not local and not (private and os.environ.get('HERMES_API_ALLOW_PRIVATE_HTTP') == '1'):
                raise ValueError('Non-loopback chat endpoints require HTTPS unless private HTTP is explicitly enabled')
        model = bounded_text(os.environ.get('HERMES_API_MODEL', 'hermes-agent'), 120)
        if not model.strip():
            raise ValueError('Chat model is not configured')
        context = (data['project'], data['slide'], data['base'])
        if session.get('context') != context:
            session['history'] = []
            session['context'] = context
        active_context = {k: view[k] for k in ('project', 'base', 'slide', 'fields')}
        prompt = ('Return JSON only: {"message":string,"ops":null or [{"slide":integer,"old":exact field old literal,"new":HTML escaped text}]}. '
                  'At most 40 ops, 4000 characters per string, message <=4000 characters. Edit only provided fields. No tools needed. Active slide: ' + json.dumps(active_context))
        messages = [{'role': 'system', 'content': prompt}] + session.get('history', []) + [{'role': 'user', 'content': data['message']}]
        encoded = json.dumps({'model': model, 'messages': messages, 'max_tokens': 3000, 'stream': False}).encode()
        if len(encoded) > LIMIT:
            raise ValueError('Conversation context too large; select another slide to reset')
        req = request.Request(base.rstrip('/') + '/chat/completions', data=encoded,
                              headers={'Authorization': 'Bearer ' + key, 'Content-Type': 'application/json'})
        # Never forward redirects with credentials.
        class NoRedirect(request.HTTPRedirectHandler):
            def redirect_request(self, *args, **kwargs):
                return None
        with request.build_opener(request.ProxyHandler({}), NoRedirect).open(req, timeout=30) as response:
            raw = response.read(LIMIT + 1)
        if len(raw) > LIMIT:
            raise ValueError('Chat response too large')
        answer = json.loads(json.loads(raw)['choices'][0]['message']['content'])
        if not isinstance(answer, dict) or set(answer) != {'message', 'ops'}:
            raise ValueError('Invalid proposal')
        bounded_text(answer['message'])
        strings = [answer['message']]
        if isinstance(answer['ops'], list):
            strings += [v for op in answer['ops'] if isinstance(op, dict) for v in op.values() if isinstance(v, str)]
        if any(key in value for value in strings):
            raise ValueError('Invalid proposal')
        if answer['ops'] is not None:
            self.validate_ops(view, answer['ops'])
        session['history'] = (session.get('history', []) + [{'role': 'user', 'content': data['message']}, {'role': 'assistant', 'content': json.dumps(answer)}])[-8:]
        return answer


def build_project(root):
    child = subprocess.Popen([sys.executable, str(Path(project.__file__)), 'build', str(root)],
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
    try:
        code = child.wait(timeout=450)
        if code:
            raise subprocess.CalledProcessError(code, 'project build')
    except subprocess.TimeoutExpired:
        os.killpg(child.pid, signal.SIGKILL)
        child.wait()
        raise


class Conflict(Exception):
    pass


class Handler(BaseHTTPRequestHandler):
    def setup(self):
        super().setup()
        self.connection.settimeout(10)

    def log_message(self, *args):
        pass  # Do not log project content, cookies, prompts or upstream errors.

    def send(self, status, body, content_type='application/json', cookie=None, attachment=None):
        if not isinstance(body, bytes):
            body = json.dumps(body).encode()
        self.send_response(status)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Length', str(len(body)))
        for k, v in {'Content-Security-Policy': "default-src 'none'; script-src 'self'; style-src 'self'; connect-src 'self'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'none'; form-action 'self'", 'X-Content-Type-Options': 'nosniff', 'Referrer-Policy': 'no-referrer', 'Cache-Control': 'no-store', 'X-Frame-Options': 'DENY', 'Permissions-Policy': 'camera=(), microphone=(), geolocation=()'}.items():
            self.send_header(k, v)
        if cookie:
            self.send_header('Set-Cookie', cookie)
        if attachment:
            self.send_header('Content-Disposition', 'attachment; filename="' + attachment + '"')
        self.end_headers()
        self.wfile.write(body)

    def session(self):
        now = time.monotonic()
        app = self.server.app
        app.sessions = {k: v for k, v in app.sessions.items() if now - v['time'] < 3600}
        cookies = http.cookies.SimpleCookie(self.headers.get('Cookie', ''))
        sid = cookies['df'].value if 'df' in cookies else ''
        if sid not in app.sessions:
            if len(app.sessions) >= 256:
                oldest = min(app.sessions, key=lambda key: app.sessions[key]['time'])
                del app.sessions[oldest]
            sid = secrets.token_urlsafe(32)
            app.sessions[sid] = {'csrf': secrets.token_urlsafe(32), 'time': now}
        session = app.sessions[sid]
        session['time'] = now
        return session, 'df=' + sid + '; HttpOnly; SameSite=Strict; Path=/' + ('; Secure' if os.environ.get('DECKFOUNDRY_SECURE_COOKIE') == '1' else '')

    def send_error(self, code, message=None, explain=None):
        self.send(code, {'error': 'Request rejected'})

    def do_GET(self):
        self.handle_request(False)

    def do_POST(self):
        self.handle_request(True)

    def handle_request(self, post):
        try:
            if not self.request_origin_ok(post):
                return self.send(403, {'error': 'Host or Origin rejected'})
            session, cookie = self.session()
            url = parse.urlsplit(self.path)
            path = url.path
            app = self.server.app
            if post:
                if self.headers.get('X-CSRF-Token') != session['csrf']:
                    return self.send(403, {'error': 'CSRF token required'})
                if self.headers.get('Content-Type') != 'application/json' or self.headers.get('Transfer-Encoding'):
                    raise ValueError('JSON required')
                length = int(self.headers.get('Content-Length', '0'))
                if not 0 < length <= (12*1024*1024 if path == '/api/create' else LIMIT):
                    raise ValueError('Payload limit exceeded')
                data = json.loads(self.rfile.read(length))
                if not isinstance(data, dict):
                    raise ValueError('Object required')
                if path == '/api/create':
                    try:
                        result = app.create(data)
                    except ValueError as exc:
                        messages = ('Use ', 'Upload ', 'HTML ', 'External ', 'Expanded ', 'Unsafe ', 'Malformed ', 'Choose ', 'Import ', 'Name ', 'This ', 'Template ')
                        message = str(exc)
                        return self.send(400, {'error': message if message.startswith(messages) else 'Malformed or unsupported upload; choose UTF-8 HTML or a standard PPTX file'})
                    except FileExistsError:
                        return self.send(400, {'error': 'This name is already in use'})
                elif path == '/api/edit':
                    try:
                        result = app.edit(data)
                    except ValueError as exc:
                        message=str(exc)
                        return self.send(400, {'error': message if message.startswith(('Keep ', 'Slide is ', 'Invalid property', 'Source limit')) else 'Edit could not be saved. Check selected elements and property limits.'})
                elif path == '/api/apply':
                    result = app.apply(data)
                elif path == '/api/chat':
                    result = app.chat(session, data)
                elif path == '/api/build':
                    if set(data) != {'project', 'base'}:
                        raise ValueError('Invalid build payload')
                    root, state, _ = app.load(data['project'])
                    if state['current'] != data['base']:
                        raise Conflict()
                    if state.get('brand'):
                        raise ValueError('Branded projects are unsupported by the generic build engine')
                    try:
                        build_project(root)
                    except (subprocess.SubprocessError, OSError):
                        return self.send(422, {'error': 'Build failed. Check that Chromium, LibreOffice and Poppler are available to the service, then retry. Custom layouts may require local review.'})
                    result = {'message': 'Build complete. Visual review required.'}
                else:
                    return self.send(404, {'error': 'Not found'})
            elif path in ('/', '/app.js', '/style.css'):
                file, mime = {'/': ('index.html', 'text/html; charset=utf-8'), '/app.js': ('app.js', 'text/javascript'), '/style.css': ('style.css', 'text/css')}[path]
                return self.send(200, (STATIC / file).read_bytes(), mime, cookie)
            elif path == '/api/projects':
                names = []
                for entry in itertools.islice(app.root.iterdir(), 1000):
                    if entry.name == '_templates': continue
                    try:
                        _, state, _ = app.load(entry.name)
                        names.append({'name': entry.name, 'current': state['current'], 'unsupported': bool(state.get('brand'))})
                    except (ValueError, OSError, KeyError, TypeError):
                        continue
                result = {'projects': names, 'csrf': session['csrf']}
            elif path == '/api/templates':
                result = {'templates': app.templates()}
            elif path == '/api/slide':
                q = parse.parse_qs(url.query, strict_parsing=True)
                result = app.view(q['project'][0], int(q['slide'][0]))
            elif path.startswith('/download/'):
                pieces = path.split('/')
                if len(pieces) != 6:
                    raise ValueError('Invalid download')
                _, _, name, version, kind, filename = pieces
                _, state, _ = app.load(name)
                if version != state['current'] or kind not in ('deck', 'render', 'html-source'):
                    raise ValueError('Only current output supported')
                if (kind == 'deck' and filename == 'deck.pptx'):
                    parts, mime = [], 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
                elif kind == 'render' and filename == 'deck.pdf':
                    parts, mime = ['render'], 'application/pdf'
                elif kind in ('render', 'html-source') and re.fullmatch(r'slide-[0-9]{1,3}\.png', filename):
                    parts, mime = [kind], 'image/png'
                else:
                    raise ValueError('Invalid artifact')
                p = app.path(name, 'versions', version, *parts, filename)
                if p.stat().st_size > 64 * 1024 * 1024:
                    raise ValueError('Artifact too large')
                return self.send(200, p.read_bytes(), mime, cookie, filename)
            else:
                return self.send(404, {'error': 'Not found'})
            self.send(200, result, cookie=cookie)
        except Conflict:
            self.send(409, {'error': 'Stale revision; reload before applying'})
        except (ValueError, KeyError, TypeError, OSError, IndexError):
            self.send(400, {'error': 'Invalid request or unavailable resource'})
        except Exception:
            self.send(502, {'error': 'Operation failed'})

    def request_origin_ok(self, post):
        host = self.headers.get('Host', '').strip().lower()
        if not host:
            return False
        allowed = allowed_origins()
        if not any(parse.urlsplit(origin).netloc == host for origin in allowed):
            return False
        origin = self.headers.get('Origin')
        if post and not origin:
            return False
        if origin:
            try:
                normalized = normalize_origin(origin)
            except ValueError:
                return False
            if normalized not in allowed or parse.urlsplit(normalized).netloc != host:
                return False
        return True


def main():
    server = HTTPServer((os.environ.get('DECKFOUNDRY_BIND', '127.0.0.1'), int(os.environ.get('DECKFOUNDRY_PORT', '8080'))), Handler)
    server.app = App(os.environ.get('DECKFOUNDRY_PROJECTS_DIR', 'projects'))
    server.serve_forever()


if __name__ == '__main__':
    main()
