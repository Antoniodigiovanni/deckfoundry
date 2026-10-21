"""Executable tests for scratch generation, template inheritance, package checks."""
import json
import subprocess
import sys
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from engine import analyze, template_generate, validate, load_spec


def fixture(tmp_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.slide_masters[0].background.fill.solid()
    prs.slide_masters[0].background.fill.fore_color.rgb = RGBColor(239, 246, 248)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Template example'
    slide.placeholders[1].text = 'Replace me'
    p = tmp_path / 'source.pptx'
    prs.save(p)
    return p


def test_analyze_and_brand(tmp_path):
    p = fixture(tmp_path)
    brand = tmp_path / 'brand'
    profile = analyze(p, tmp_path / 'profile.json', brand)
    assert profile['size_inches'][0] > 13
    assert len(profile['layouts']) >= 2
    assert len(profile['slides']) == 1
    assert (brand / 'template.pptx').exists()
    assert (brand / 'profile.json').exists()
    assert json.loads((brand / 'brand.json').read_text())['default_layout'] == 1


def generic_spec(tmp_path):
    path = tmp_path / 'spec.json'
    path.write_text(json.dumps({'title': 'Generic deck', 'slides': [{'title': f'Slide {i}', 'bullets': ['Option'], 'notes': 'Notes'} for i in range(5)]}))
    return path


def test_template_five_editable(tmp_path):
    p = fixture(tmp_path)
    brand = tmp_path / 'brand'
    analyze(p, tmp_path / 'profile.json', brand)
    cfg = json.loads((brand / 'brand.json').read_text())
    cfg['default_layout'] = 1
    (brand / 'brand.json').write_text(json.dumps(cfg))
    out = tmp_path / 'template-output.pptx'
    template_generate(generic_spec(tmp_path), brand, out)
    prs = Presentation(out)
    assert len(prs.slides) == 5
    assert len(prs.slide_masters) == 1
    assert all(any(sh.has_text_frame and 'Option' in sh.text for sh in sl.shapes) for sl in [prs.slides[2]])
    assert all(sl.slide_layout.name == prs.slide_layouts[1].name for sl in prs.slides)
    assert validate(out)['passed']
    assert all(len(sl.shapes) > 1 for sl in prs.slides)


def test_scratch_five_editable(tmp_path):
    out = tmp_path / 'scratch.pptx'
    subprocess.run(['node', str(ROOT / 'src/generate.js'), str(generic_spec(tmp_path)), str(out)], check=True)
    prs = Presentation(out)
    assert len(prs.slides) == 5
    assert all(any(sh.has_text_frame and sh.text for sh in sl.shapes) for sl in prs.slides)
    assert validate(out)['passed']


def test_reject_invalid_spec(tmp_path):
    bad = tmp_path / 'bad.json'
    bad.write_text('{"slides":[{"title":""}]}')
    try:
        load_spec(bad)
    except ValueError:
        pass
    else:
        raise AssertionError('invalid spec accepted')
