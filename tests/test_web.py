"""Generic temporary projects and in-memory HTTP requests/upstream mocks."""
import io
import email
from types import SimpleNamespace
import importlib.util
import json
from pathlib import Path
import sys
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
import project
spec = importlib.util.spec_from_file_location('deckweb', ROOT / 'web/app.py')
web = importlib.util.module_from_spec(spec)
spec.loader.exec_module(web)
SOURCE = '<html><body><section class="slide"><h1>Title</h1><p>Body &amp; text</p><ul><li>Bullet</li></ul><aside class="speaker-notes">Notes</aside></section><section class="slide"><h1>Other slide</h1></section></body></html>'


@pytest.fixture
def app(tmp_path):
    source = tmp_path / 'input.html'
    source.write_text(SOURCE)
    root = tmp_path / 'projects'
    root.mkdir()
    project.init(root / 'generic', source)
    return web.App(root)



@pytest.fixture
def client(app):
    def call(path, data=None, cookie=None, csrf=None, raw=None, host='localhost:8080', origin='http://localhost:8080'):
        headers = {'Host': host}
        if origin:
            headers['Origin'] = origin
        if cookie:
            headers['Cookie'] = cookie
        if csrf:
            headers['X-CSRF-Token'] = csrf
        body = raw if raw is not None else json.dumps(data).encode() if data is not None else None
        if body is not None:
            headers['Content-Type'] = 'application/json'
            headers['Content-Length'] = str(len(body))
        wire = ((('POST' if body is not None else 'GET') + ' ' + path + ' HTTP/1.0\r\n') + ''.join(k + ': ' + v + '\r\n' for k, v in headers.items()) + '\r\n').encode() + (body or b'')
        handler = web.Handler.__new__(web.Handler)
        handler.rfile = io.BytesIO(wire)
        handler.wfile = io.BytesIO()
        handler.server = SimpleNamespace(app=app)
        handler.client_address = ('127.0.0.1', 1234)
        handler.handle_one_request()
        head, payload = handler.wfile.getvalue().split(b'\r\n\r\n', 1)
        lines = head.decode().split('\r\n', 1)
        return int(lines[0].split()[1]), email.message_from_string(lines[1]), payload
    call.app = app
    return call


def session(client):
    code, headers, raw = client('/api/projects')
    assert code == 200
    return headers['Set-Cookie'].split(';')[0], json.loads(raw)['csrf']


def patch():
    return {'project': 'generic', 'slide': 1, 'base': 'v0001', 'ops': [{'slide': 1, 'old': 'Title', 'new': 'Updated &lt;safe&gt;'}]}


def test_list_view_headers(client):
    code, headers, raw = client('/api/projects')
    assert json.loads(raw)['projects'] == [{'name': 'generic', 'current': 'v0001', 'unsupported': False}]
    assert headers['X-Content-Type-Options'] == 'nosniff'
    assert headers['X-Frame-Options'] == 'DENY'
    assert "default-src 'none'" in headers['Content-Security-Policy']
    assert not headers.get('Access-Control-Allow-Origin')
    assert 'HttpOnly' in headers['Set-Cookie'] and 'SameSite=Strict' in headers['Set-Cookie']
    code, _, raw = client('/api/slide?project=generic&slide=1')
    data = json.loads(raw)
    assert [f['kind'] for f in data['fields']] == ['title', 'text', 'bullet', 'notes']
    assert data['fields'][1]['text'] == 'Body & text'
    assert 'html' not in data
    assert client('/')[0] == 200


def test_host_and_origin_rebinding_rejected(client):
    assert client('/api/projects', host='attacker.example', origin='http://attacker.example')[0] == 403
    assert client.app.sessions == {}
    cookie, csrf = session(client)
    assert client('/api/apply', patch(), cookie, csrf, host='localhost:8080', origin='http://attacker.example')[0] == 403
    assert client('/api/apply', patch(), cookie, csrf, origin=None)[0] == 403


def test_session_capacity_evicts_oldest_instead_of_locking_out(app, client):
    now = web.time.monotonic()
    app.sessions = {f'session-{i}': {'csrf': f'token-{i}', 'time': now - 300 + i}
                    for i in range(256)}
    code, headers, raw = client('/api/projects')
    assert code == 200
    assert json.loads(raw)['projects'][0]['name'] == 'generic'
    new_sid = headers['Set-Cookie'].split(';')[0].removeprefix('df=')
    assert new_sid in app.sessions
    assert len(app.sessions) == 256
    assert 'session-0' not in app.sessions


def test_apply_immutable_stale(client, app):
    cookie, csrf = session(client)
    assert client('/api/apply', patch(), cookie, csrf)[0] == 200
    assert (app.root / 'generic/versions/v0001/source.html').read_text() == SOURCE
    assert 'Updated &lt;safe&gt;' in (app.root / 'generic/versions/v0002/source.html').read_text()
    assert client('/api/apply', patch(), cookie, csrf)[0] == 409
    assert not (app.root / 'generic/versions/v0003').exists()


@pytest.mark.parametrize('change', [
    {'ops': []}, {'ops': None}, {'slide': True}, {'slide': 0}, {'extra': 1},
    {'ops': [{'slide': 1, 'old': 'Title', 'new': '<script>alert(1)</script>'}]},
    {'ops': [{'slide': 2, 'old': 'Other slide', 'new': 'Bad'}]},
    {'ops': [{'slide': 1, 'old': 'class="slide"', 'new': 'Bad'}]},
    {'ops': [{'slide': 1, 'old': 'Title', 'new': 'x' * 4001}]},
    {'ops': [{'slide': 1, 'old': [], 'new': 'x'}]},
])
def test_validation(client, app, change):
    cookie, csrf = session(client)
    data = patch(); data.update(change)
    assert client('/api/apply', data, cookie, csrf)[0] == 400
    assert project.revision(app.root / 'generic')[0]['current'] == 'v0001'


def test_csrf_payload_and_errors(client):
    cookie, csrf = session(client)
    assert client('/api/apply', patch(), cookie)[0] == 403
    assert client('/api/apply', cookie=cookie, csrf=csrf, raw=b'{' )[0] == 400
    assert client('/api/apply', cookie=cookie, csrf=csrf, raw=b'x' * 65537)[0] == 400
    code, headers, _ = client('/missing')
    assert code == 404 and headers['Cache-Control'] == 'no-store'


def test_paths_and_downloads(client, app, tmp_path):
    for name in ('..', '%2e%2e', 'generic%2f..', '%2Ftmp', 'generic%5c..'):
        assert client('/api/slide?project=' + name + '&slide=1')[0] == 400
    outside = tmp_path / 'outside'; outside.mkdir()
    (app.root / 'escape').symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError): app.load('escape')
    (app.root / 'generic/versions/v0001/deck.pptx').symlink_to(outside / 'secret')
    assert client('/download/generic/v0001/deck/deck.pptx')[0] == 400
    (app.root / 'generic/versions/v0001/deck.pptx').unlink()
    (app.root / 'generic/versions/v0001/deck.pptx').write_bytes(b'generic artifact')
    code, headers, body = client('/download/generic/v0001/deck/deck.pptx')
    assert code == 200 and body == b'generic artifact'
    assert 'attachment' in headers['Content-Disposition']
    assert client('/download/generic/v0000/deck/deck.pptx')[0] == 400
    assert client('/download/generic/v0001/deck/source.html')[0] == 400
    (app.root / 'generic/versions/v0001/source.html').unlink()
    (app.root / 'generic/versions/v0001/source.html').symlink_to(outside / 'secret')
    with pytest.raises(ValueError): app.load('generic')


def test_build_guard_and_call(client, app, monkeypatch):
    cookie, csrf = session(client)
    calls = []
    monkeypatch.setattr(web, 'build_project', lambda root: calls.append(root))
    data = {'project': 'generic', 'base': 'v0001'}
    assert client('/api/build', data, cookie, csrf)[0] == 200
    assert calls[0] == app.root / 'generic'
    state = app.root / 'generic/project.json'
    value = json.loads(state.read_text()); value['brand'] = 'local'; state.write_text(json.dumps(value))
    assert client('/api/build', data, cookie, csrf)[0] == 400
    assert len(calls) == 1


def test_chat_sessions_key_proposals(client, app, monkeypatch):
    calls = []
    key = 'mock-server-only-key'
    class Mock:
        def open(self, req, timeout):
            assert timeout == 30
            body = json.loads(req.data)
            calls.append((web.parse.urlsplit(req.full_url).path, req.headers['Authorization'], body))
            answer = {'message': 'Review this', 'ops': patch()['ops']}
            content = json.dumps({'choices': [{'message': {'content': json.dumps(answer)}}]}).encode()
            return io.BytesIO(content)
    monkeypatch.setattr(web.request, 'build_opener', lambda *a: Mock())
    base = 'http://127.0.0.1:8642'
    monkeypatch.setenv('HERMES_API_BASE_URL', base + '/v1')
    monkeypatch.setenv('HERMES_API_KEY', key)
    cookie, csrf = session(client)
    data = {'project': 'generic', 'slide': 1, 'base': 'v0001', 'message': 'Improve this'}
    code, _, raw = client('/api/chat', data, cookie, csrf)
    assert code == 200 and key.encode() not in raw
    assert project.revision(app.root / 'generic')[0]['current'] == 'v0001'
    assert calls[-1][0] == '/v1/chat/completions'
    assert calls[-1][1] == 'Bearer ' + key
    assert calls[-1][2]['model'] == 'hermes-agent'
    messages = calls[-1][2]['messages']
    assert len(messages) == 2 and 'Other slide' not in json.dumps(messages)
    assert client('/api/chat', data, cookie, csrf)[0] == 200
    assert len(calls[-1][2]['messages']) == 4
    cookie2, csrf2 = session(client)
    assert cookie2 != cookie
    assert client('/api/chat', data, cookie2, csrf2)[0] == 200
    assert len(calls[-1][2]['messages']) == 2
    data['slide'] = 2
    # Mock proposal for slide 1 is rejected against active slide 2.
    assert client('/api/chat', data, cookie, csrf)[0] == 400
    assert len(calls[-1][2]['messages']) == 2
    assert client('/api/apply', patch(), cookie, csrf)[0] == 200


@pytest.mark.parametrize('content', [
    b'x' * 65537,
    b'not json',
    json.dumps({'choices': [{'message': {'content': json.dumps({'message': 'private-key', 'ops': None})}}]}).encode(),
    json.dumps({'choices': [{'message': {'content': json.dumps({'message': 'bad', 'ops': [{'slide': 1, 'old': 'Title', 'new': '<img src=x>'}]})}}]}).encode(),
])
def test_chat_rejects_unsafe_response(client, monkeypatch, content):
    class Mock:
        def open(self, req, timeout): return io.BytesIO(content)
    monkeypatch.setattr(web.request, 'build_opener', lambda *a: Mock())
    monkeypatch.setenv('HERMES_API_BASE_URL', 'http://localhost:8642/v1')
    monkeypatch.setenv('HERMES_API_KEY', 'private-key')
    cookie, csrf = session(client)
    code, _, raw = client('/api/chat', {'project': 'generic', 'slide': 1, 'base': 'v0001', 'message': 'Help'}, cookie, csrf)
    assert code == 400 and b'private-key' not in raw


def test_symlink_root_and_metadata(app, tmp_path):
    alias = tmp_path / 'alias'
    alias.symlink_to(app.root, target_is_directory=True)
    with pytest.raises(ValueError): web.App(alias)
    state = app.root / 'generic/project.json'
    data = json.loads(state.read_text()); data['current'] = '../../elsewhere'
    state.write_text(json.dumps(data))
    with pytest.raises(ValueError): app.load('generic')


def test_build_deadline_kills_group(monkeypatch, tmp_path):
    killed = []
    class Child:
        pid = 12345
        def wait(self, timeout=None):
            if timeout: raise web.subprocess.TimeoutExpired('build', timeout)
            return -9
    monkeypatch.setattr(web.subprocess, 'Popen', lambda *a, **kw: Child())
    monkeypatch.setattr(web.os, 'killpg', lambda pid, sig: killed.append((pid, sig)))
    with pytest.raises(web.subprocess.TimeoutExpired): web.build_project(tmp_path)
    assert killed == [(12345, web.signal.SIGKILL)]


def test_http_chat_refuses_remote_endpoint(app, monkeypatch):
    calls = []
    monkeypatch.setenv('HERMES_API_BASE_URL', 'http://api.example.test/v1')
    monkeypatch.setenv('HERMES_API_KEY', 'server-only-key')
    monkeypatch.setattr(web.request, 'build_opener', lambda *a: calls.append(a))
    with pytest.raises(ValueError):
        app.chat({'history': []}, {'project': 'generic', 'slide': 1, 'base': 'v0001', 'message': 'Help'})
    assert calls == []


def test_private_hermes_bridge_requires_explicit_opt_in(app, monkeypatch):
    calls = []

    class Mock:
        def open(self, req, timeout):
            calls.append(req.full_url)
            payload = {'choices': [{'message': {'content': json.dumps({'message': 'Ready', 'ops': None})}}]}
            return io.BytesIO(json.dumps(payload).encode())

    monkeypatch.setattr(web.request, 'build_opener', lambda *a: Mock())
    monkeypatch.setenv('HERMES_API_BASE_URL', 'http://172.17.0.1:8642/v1')
    monkeypatch.setenv('HERMES_API_KEY', 'mock-server-only-key')
    data = {'project': 'generic', 'slide': 1, 'base': 'v0001', 'message': 'Hello'}
    with pytest.raises(ValueError):
        app.chat({'history': []}, data)
    assert calls == []

    monkeypatch.setenv('HERMES_API_ALLOW_PRIVATE_HTTP', '1')
    answer = app.chat({'history': []}, data)
    assert answer == {'message': 'Ready', 'ops': None}
    assert calls == ['http://172.17.0.1:8642/v1/chat/completions']


def upload(name='imported', filename='input.html', raw=SOURCE.encode(), kind='project'):
    return {'name': name, 'filename': filename, 'content': web.base64.b64encode(raw).decode(), 'kind': kind}


def test_create_template_select_and_private_original(client, app):
    cookie, csrf = session(client)
    assert json.loads(client('/api/templates')[2]) == {'templates': []}
    assert client('/api/create', upload('theme', kind='template'), cookie, csrf)[0] == 200
    listing = json.loads(client('/api/templates')[2])['templates']
    assert listing == [{'name': 'theme.html'}]
    assert client('/api/create', {'name':'from-template','template':'theme.html'}, cookie, csrf)[0] == 200
    assert app.view('from-template', 1)['fields'][0]['text'] == 'Title'
    names = [p['name'] for p in json.loads(client('/api/projects')[2])['projects']]
    assert '_templates' not in names and 'theme' not in names
    assert client('/api/create', upload(), cookie, csrf)[0] == 200
    assert (app.root/'imported/original.html').read_bytes() == SOURCE.encode()
    assert client('/api/create', upload(), cookie, csrf)[0] == 400
    assert client('/api/create', {'name':'blank'}, cookie, csrf)[0] == 200


@pytest.mark.parametrize('path,data', [
    ('/api/create', {'name':'new'}),
    ('/api/edit', {'project':'generic','base':'v0001','slide':1,'action':'add'}),
])
def test_new_routes_require_origin_session_csrf(client, path, data):
    cookie, csrf = session(client)
    assert client(path, data)[0] == 403
    assert client(path, data, cookie, csrf, origin=None)[0] == 403
    assert client(path, data, cookie, csrf, origin='https://evil.example')[0] == 403
    assert client(path, data, cookie, csrf)[0] == 200
    assert client('/api/templates', host='evil.example')[0] == 403


@pytest.mark.parametrize('filename,raw', [
    ('../escape.html', SOURCE.encode()), ('x.svg', b'<svg/>'),
    ('x.html', b'<html>no slides</html>'), ('x.pptx', b'not a package'),
    ('x.html', b'\xff'), ('x.html', b'x'*(web.SOURCE_LIMIT+1)),
    ('x.html', b'<!DOCTYPE html SYSTEM "file:///etc/passwd"><section class="slide"><p>x</p></section>'),
])
def test_import_rejections(client, filename, raw):
    cookie, csrf = session(client)
    assert client('/api/create', upload(filename=filename, raw=raw), cookie, csrf)[0] == 400
    assert not (client.app.root/'imported').exists()


def test_import_strips_active_content_and_remote_resources(app):
    raw=b'<section class="slide" onclick="alert(1)"><h1 style="color:#112233;background:url(https://evil.example)">Safe<script>bad()</script></h1><iframe src="file:///secret"></iframe><img src="https://evil.example/a"><p onmouseover="bad()">Text</p></section>'
    text=web.import_source('safe.html', raw)
    assert all(x not in text for x in ('script','iframe','evil.example','onclick','onmouseover'))
    assert 'color:#112233' in text and 'Safe' in text


def test_pptx_sequence_geometry_images_and_bad_zip(app):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    deck=Presentation()
    for title in ('Opening','Closing'):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        shape.text=title
    picture=io.BytesIO();Image.new('RGB',(10,10),'red').save(picture,format='PNG');picture.seek(0)
    deck.slides[0].shapes.add_picture(picture, Inches(2), Inches(3), width=Inches(1))
    raw=io.BytesIO();deck.save(raw)
    app.create(upload(raw=raw.getvalue(),filename='deck.pptx'))
    view=app.view('imported',1)
    assert view['fields'][0]['text']=='Opening'
    assert view['fields'][0]['style']['left'].endswith('px')
    assert view['images'][0]['src'].startswith('data:image/png;base64,')
    assert app.view('imported',2)['fields'][0]['text']=='Closing'
    bad=io.BytesIO()
    with web.zipfile.ZipFile(bad,'w') as archive:
        archive.writestr('../escape.xml','<x/>')
    with pytest.raises(ValueError): web.import_source('bad.pptx',bad.getvalue())


def test_format_batch_atomic_stale_and_slide_actions(client,app):
    cookie,csrf=session(client)
    data={'project':'generic','base':'v0001','slide':1,'action':'format','edits':[
        {'field':0,'text':'Edited <safe>','styles':{'font-size':'60px','color':'#123456','text-align':'center'}},
        {'styles':{'background':'#abcdef'}}]}
    assert client('/api/edit',data,cookie,csrf)[0]==200
    view=app.view('generic',1)
    assert view['base']=='v0002' and view['background']=='#abcdef'
    assert view['fields'][0]['text']=='Edited <safe>'
    assert view['fields'][0]['style']['font-size']=='60px'
    assert client('/api/edit',data,cookie,csrf)[0]==409
    data['base']='v0002';data['edits'][1]={'field':500,'text':'bad'}
    assert client('/api/edit',data,cookie,csrf)[0]==400
    assert app.view('generic',1)['base']=='v0002'
    for action,count in [('duplicate',3),('down',3),('up',3),('delete',2),('add',3)]:
        data={'project':'generic','base':app.view('generic',1)['base'],'slide':2 if action=='up' else 1,'action':action}
        assert client('/api/edit',data,cookie,csrf)[0]==200
        assert len(app.view('generic',1)['slides'])==count
    assert (app.root/'generic/versions/v0001/source.html').read_text()==SOURCE


def test_template_symlinks_and_selection_escape(app,tmp_path):
    folder=app.root/'_templates';folder.mkdir()
    (folder/'escape.html').symlink_to(tmp_path/'outside')
    assert app.templates()==[]
    for name in ('../outside.html','escape.html'):
        with pytest.raises(ValueError):app.create({'name':'new','template':name})


def test_format_rejects_css_and_markup(client):
    cookie,csrf=session(client)
    for styles in ({'color':'url(https://evil.example)'},{'font-size':'999px'},{'position':'fixed'}):
        data={'project':'generic','base':'v0001','slide':1,'action':'format','field':0,'styles':styles}
        assert client('/api/edit',data,cookie,csrf)[0]==400
