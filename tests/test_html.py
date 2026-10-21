import json
import subprocess
import sys
from pathlib import Path
import zipfile
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]


def test_html_editable_offline(tmp_path):
    src = tmp_path / 'input.html'
    src.write_text('''<html><head><style>.stage{width:1440px;height:810px}.slide{position:absolute;inset:0;background:#06142c;color:white}h1{font:40px Arial}</style></head><body><div class="stage"><section class="slide"><h1>Editable title</h1><svg width="300" height="120"><rect x="10" y="10" width="100" height="70" fill="#20e0c0"/><text x="20" y="40" font-size="15">Node label</text><path d="M110 40 C 130 40 150 60 190 60" fill="none" stroke="#20e0c0"/></svg></section></div><script>throw Error("Must never run")</script><img src="https://example.invalid/private"></body></html>''')
    out = tmp_path / 'result.pptx'
    subprocess.run([sys.executable, str(ROOT / 'src/engine.py'), 'html', str(src), '--out', str(out), '--slides', '1'], check=True, timeout=90)
    prs = Presentation(out)
    assert len(prs.slides) == 1
    texts = [s.text for s in prs.slides[0].shapes if s.has_text_frame]
    assert 'Editable title' in texts
    assert 'Node label' in texts
    assert len(prs.slides[0].shapes) > 5
    with zipfile.ZipFile(out) as z:
        assert not any(n.startswith('ppt/media/') and not n.endswith('/') for n in z.namelist())
    spec = json.loads((tmp_path / 'html-source/spec.json').read_text())
    assert spec['network'] == 'denied'
    assert (tmp_path / 'html-source/slide-1.png').exists()


def test_html_invalid_slide_count(tmp_path):
    src = tmp_path / 'empty.html'
    src.write_text('<html><body></body></html>')
    proc = subprocess.run(['node', str(ROOT / 'src/html-convert.js'), str(src), str(tmp_path / 'bad.pptx'), '5'], capture_output=True, text=True, timeout=60)
    assert proc.returncode != 0
    assert not (tmp_path / 'bad.pptx').exists()


def test_studio_frontend_state_and_safe_canvas():
    """Exercise real frontend functions without requiring browser startup."""
    script = r'''
const fs=require('fs'),vm=require('vm'),assert=require('assert');
class Element {
 constructor(){this.children=[];this.style={};this.dataset={};this.value='';this.hidden=false;this.classList={add(){},remove(){},toggle(){}};}
 append(...items){this.children.push(...items);} replaceChildren(...items){this.children=items;}
 setAttribute(k,v){this[k]=v;} focus(){} scrollIntoView(){} querySelectorAll(){return [];} querySelector(selector){const match=selector.match(/data-field="(\d+)"/);return match?this.children.find(e=>String(e.dataset.field)===match[1]):new Element();}
}
const elements=new Map();const document={getElementById(id){if(!elements.has(id))elements.set(id,new Element());return elements.get(id);},createElement(){return new Element();},body:new Element(),addEventListener(){}};
const context={document,window:{addEventListener(){}},Option:function(t,v){this.text=t;this.value=v;},URLSearchParams,console};
vm.createContext(context);
let source=fs.readFileSync('web/app.js','utf8');source=source.replace(/run\(async\(\)=>\{await projects\(\);status\('Ready'\);\}\);\s*$/,'');
vm.runInContext(source,context);
vm.runInContext(`
 active={project:'test',slide:1,base:'v0001',background:'#ffffff',images:[],fields:[{kind:'title',text:'<img onerror=bad()>',old:'Title',style:{}}]};
 renderCanvas(active,$('preview'),true);
 if($('preview').children[0].textContent!=='<img onerror=bad()>')throw Error('Unsafe rendering');
 selected=0;$('content').value='Changed';$('content').oninput();
 if(!dirty||$('preview').children[0].textContent!=='Changed')throw Error('Edit did not update canvas');
 pending=[{old:'x'}];$('review').hidden=false;mark();
 if(pending!==null||!$('review').hidden)throw Error('Local edit retained stale proposal');
 `,context);
(async()=>{
 await vm.runInContext(`(async()=>{
 let calls=[];api=async(path,data)=>{calls.push({path,data});return {current:'v0002'};};
 await save();
 if(active.base!=='v0002'||dirty||calls.length!==1||calls[0].data.edits[0].text!=='Changed')throw Error('Save transaction failed');
 active.fields[0].text='Pending';mark();api=async()=>{throw Error('Stale revision');};
 await run(save);
 if(!dirty||active.fields[0].text!=='Pending'||!$('status').textContent.includes('Stale'))throw Error('Failed write lost edits');
 pending=[{old:'x'}];$('rejectProposal').onclick();if(pending!==null)throw Error('Reject failed');
 dirty=false;active.fields[0].original=active.fields[0].text;let applied=false;
 api=async(path,data)=>{
  if(path==='/api/chat')return {message:'Review',ops:[{slide:1,old:'Title',new:'Approved'}]};
  if(path==='/api/apply'){applied=true;return {current:'v0003'};}
  if(path.startsWith('/api/slide'))return {project:'test',slide:1,base:'v0003',background:'#ffffff',images:[],slides:['Approved'],fields:[{kind:'title',text:'Approved',old:'Approved',style:{}}]};
  throw Error('Unexpected route');
 };
 $('prompt').value='Improve title';await chat();
 if(!pending||applied||$('review').hidden)throw Error('Proposal applied without review');
 await $('applyProposal').onclick();
 if(!applied||active.base!=='v0003'||$('preview').children[0].textContent!=='Approved'||pending!==null)throw Error('Proposal did not update canvas');
 })()`,context);
})().catch(e=>{console.error(e);process.exitCode=1;});
'''
    subprocess.run(['node', '-e', script], cwd=ROOT, check=True, timeout=15)
